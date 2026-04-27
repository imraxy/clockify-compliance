from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook

from app.models import AttendanceDay, TimeEntry
from app.services.csv_import import (
    import_attendance_excel,
    import_time_entries_csv,
    import_time_entries_excel,
)


def test_import_time_entries_csv_roundtrip(db_session, seed_users):
    emp = seed_users["employee"]
    csv_text = (
        "employee_email,start,end,description,project\n"
        f"{emp.email},2026-01-10T09:00:00+00:00,2026-01-10T17:00:00+00:00,dev,projA\n"
    )
    res = import_time_entries_csv(db_session, csv_text)
    assert res["imported"] == 1
    assert not res["errors"]


def _workbook_bytes(wb: Workbook) -> bytes:
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def test_import_time_entries_excel_imports_and_deduplicates_rows(db_session, seed_users):
    emp = seed_users["employee"]
    wb = Workbook()
    ws = wb.active
    ws.append(["employee_email", "start", "end", "description", "project"])
    row = [
        emp.email,
        "2026-01-10T09:00:00+00:00",
        "2026-01-10T17:00:00+00:00",
        "dev",
        "projA",
    ]
    ws.append(row)
    ws.append(row)

    res = import_time_entries_excel(db_session, _workbook_bytes(wb))

    assert res["imported"] == 2
    assert not res["errors"]
    entries = db_session.query(TimeEntry).all()
    assert len(entries) == 1
    assert entries[0].project_name == "projA"


def test_import_time_entries_excel_reports_unknown_user(db_session):
    wb = Workbook()
    ws = wb.active
    ws.append(["employee_email", "start", "end", "description", "project"])
    ws.append([
        "missing@example.com",
        "2026-01-10T09:00:00+00:00",
        "2026-01-10T17:00:00+00:00",
        "dev",
        "projA",
    ])

    res = import_time_entries_excel(db_session, _workbook_bytes(wb))

    assert res["imported"] == 0
    assert res["errors"] == ["row 2: unknown user missing@example.com"]


def test_import_attendance_excel_returns_imported_and_deduplicates_days(db_session, seed_users):
    emp = seed_users["employee"]
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value=2026)
    ws.cell(row=3, column=4, value="Email Id")
    ws.cell(row=3, column=7, value=datetime(2026, 1, 10))
    ws.cell(row=3, column=8, value=datetime(2026, 1, 10))
    ws.cell(row=4, column=4, value=emp.email)
    ws.cell(row=4, column=7, value="P")
    ws.cell(row=4, column=8, value="WO")

    res = import_attendance_excel(db_session, _workbook_bytes(wb))

    assert res["imported"] == 1
    assert res["rows"] == 1
    assert not res["errors"]
    rows = db_session.query(AttendanceDay).all()
    assert len(rows) == 1
    assert rows[0].code == "P"


def test_import_attendance_excel_reports_unknown_user(db_session):
    wb = Workbook()
    ws = wb.active
    ws.cell(row=3, column=4, value="Email Id")
    ws.cell(row=3, column=7, value=datetime(2026, 1, 10))
    ws.cell(row=4, column=4, value="missing@example.com")
    ws.cell(row=4, column=7, value="P")

    res = import_attendance_excel(db_session, _workbook_bytes(wb))

    assert res["imported"] == 0
    assert res["rows"] == 0
    assert res["errors"] == ["row 4: unknown user missing@example.com"]
